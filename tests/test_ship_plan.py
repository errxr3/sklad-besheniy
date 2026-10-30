"""Экран «План отгрузки»: заказы дня против остатка, задач и погруженного (2026-10-30)."""

from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.services.batches import receive_batch, receive_demand
from app.services.ship_plan import build_ship_plan, visible_rows
from app.services.shipping import join_channel, scan_loading
from app.services.tasks import create_task
from app.services.units import accept_unit, create_printed_units
from app.statuses import ROLE_ASSEMBLER, ROLE_LOADER, ROLE_SHIFT_LEAD
from app.timeutil import today_msk
from tests.helpers import login, make_article, make_user, site_by_name

KOM = "Коминтерна"


def _line(posting: str, channel: str, article: str, qty: int = 1) -> dict:
    market = {"Ozon": "ozon", "КГТ": "ozon", "СДЭК": "cdek"}.get(channel, "wb")
    return {"market": market, "channel": channel, "posting": posting, "article": article, "qty": qty,
            "production": KOM}


def _world(db: Session):  # type: ignore[no-untyped-def]
    site = site_by_name(db, KOM)
    worker = make_user(db, name="Иван План", role=ROLE_ASSEMBLER, home_site_id=site.id)
    lead = make_user(db, name="Начальник План", role=ROLE_SHIFT_LEAD, home_site_id=site.id)
    loader = make_user(db, name="Грузчик План", role=ROLE_LOADER, home_site_id=site.id)
    a = make_article(db, name="Zerkalo_50x130", sort_order=1)
    b = make_article(db, name="Zerkalo_60x160", sort_order=2)
    make_article(db, name="Zerkalo_70x170", sort_order=3)  # без заказов — в плане не показывается
    units = create_printed_units(db, article_id=a.id, qty=3, site_id=site.id, assembler_id=worker.id,
                                 actor_id=worker.id)
    db.commit()
    for unit in units:
        accept_unit(db, code=unit.number, actor=worker)
    create_task(db, actor=lead, day=today_msk(), site_id=site.id, user_id=worker.id, kind="article",
                article_id=b.id, qty_plan=2)
    receive_demand(db, day=today_msk(), lines=[
        _line("OZ-1", "Ozon", "Zerkalo_50x130"),
        _line("OZ-2", "Ozon", "Zerkalo_50x130"),
        _line("WB-1", "WB", "Zerkalo_50x130"),
        _line("WB-2", "WB", "Zerkalo_60x160"),
        _line("OZ-3", "Ozon", "Zerkalo_60x160", qty=3),  # 3 штуки, а в работе 2 — не поедет
        _line("OZ-9", "Ozon", "Неизвестный_артикул"),
    ])
    db.commit()
    return site, lead, loader, units


def test_plan_counts_need_decisions_and_loaded(db: Session) -> None:
    site, _, loader, units = _world(db)
    receive_batch(db, day=today_msk(), market="ozon", channel="Ozon", truck="Озон1", items=[
        {"posting": "OZ-1", "article": "Zerkalo_50x130", "production": KOM, "codes": ["%OZ1"]},
    ])
    db.commit()
    pallet = join_channel(db, user=loader, site=site, channel="Ozon", truck="Озон1")
    db.commit()
    assert scan_loading(db, code=units[0].number, pallet=pallet, actor=loader, ship_code="%OZ1").ok
    db.commit()

    plan = build_ship_plan(db, day=today_msk())
    rows = {row.name: row for row in visible_rows(plan, show_all=False) if not row.is_header}
    assert set(rows) == {"Zerkalo_50x130", "Zerkalo_60x160"}
    ozon, wb = plan.columns.index("Озон"), plan.columns.index("WB")
    first = rows["Zerkalo_50x130"]
    assert first.need[ozon] == 2 and first.need[wb] == 1 and first.loaded[ozon] == 1
    assert first.go == 3 and first.state == "ok"
    second = rows["Zerkalo_60x160"]
    assert second.in_work == 2 and second.no == 3 and second.state == "no"
    assert plan.unknown == ["Неизвестный_артикул"]


def test_plan_page_and_xlsx(client, db: Session) -> None:  # type: ignore[no-untyped-def]
    _, lead, _, _ = _world(db)
    login(client, lead)
    page = client.get("/shipments/plan").text
    assert "План отгрузки" in page and "Zerkalo_60x160" in page and "не хватит" in page
    assert "Zerkalo_70x170" not in page
    assert "Zerkalo_70x170" in client.get("/shipments/plan?full=1").text
    response = client.get("/shipments/plan.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    names = [row[0] for row in sheet.iter_rows(values_only=True)]
    assert "Zerkalo_50x130" in names and names[-1] == "Всего"
