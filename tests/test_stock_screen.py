"""SPEC §12 п. 19 и 20: экран «Остатки» — порядок строк, TSV, задачи, выгрузка."""

from datetime import timedelta
from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.importer import import_top, parse_top
from app.services.stock import build_stock_table, sheet_rows, table_to_tsv
from app.services.tasks import create_task
from app.services.units import accept_unit, create_printed_units
from app.statuses import ROLE_ADMIN, ROLE_ASSEMBLER, ROLE_SHIFT_LEAD
from app.timeutil import today_msk
from tests.helpers import login, make_user, site_by_name

TOP_PATH = Path(__file__).resolve().parent.parent / "docs" / "input" / "TOP.xlsx"


def _with_catalog(db: Session) -> None:
    """Загрузить настоящий справочник."""
    import_top(db, TOP_PATH)
    db.commit()


def _make_stock(db: Session, article_name_norm: str, qty: int, site_name: str = "Коминтерна") -> None:
    """Положить на склад несколько зеркал артикула."""
    from sqlalchemy import select

    from app.models import Article

    article = db.scalar(select(Article).where(Article.name_norm == article_name_norm))
    assert article is not None
    assembler = make_user(
        db, name=f"Сборщик {article_name_norm[:10]}{site_name}{qty}", role=ROLE_ASSEMBLER
    )
    site = site_by_name(db, site_name)
    units = create_printed_units(
        db,
        article_id=article.id,
        qty=qty,
        site_id=site.id,
        assembler_id=assembler.id,
        actor_id=assembler.id,
    )
    db.commit()
    for unit in units:
        accept_unit(db, code=unit.number, actor=assembler)
    db.commit()


def test_row_order_matches_top_file(db: Session) -> None:
    """Порядок строк совпадает с порядком в TOP.xlsx, включая заголовки (SPEC §12 п. 19)."""
    _with_catalog(db)
    parsed, _ = parse_top(TOP_PATH)
    table = build_stock_table(db)

    assert table.row_count == len(parsed) == 306
    assert [row.name for row in table.rows] == [row.name for row in parsed]
    assert [row.is_header for row in table.rows] == [row.kind == "header" for row in parsed]


def test_headers_are_rows_without_numbers(db: Session) -> None:
    """Строки-заголовки подгрупп выводятся отдельными строками без чисел."""
    _with_catalog(db)
    table = build_stock_table(db)
    headers = [row for row in table.rows if row.is_header]
    assert len(headers) == 10
    assert headers[0].name == "ПОЛОТНА"
    assert headers[0].stock_text(0) == ""
    assert headers[0].tasks_text(0) == ""


def test_article_names_keep_original_spelling(db: Session) -> None:
    """Артикул выводится в исходном написании из TOP (с кириллической «х»)."""
    _with_catalog(db)
    table = build_stock_table(db)
    names = [row.name for row in table.rows]
    assert "Zerkalo_100х200" in names  # кириллическая «х», как в файле


def test_inactive_articles_keep_their_place(db: Session) -> None:
    """Отключённый артикул остаётся на месте с остатком 0 (SPEC §6.9)."""
    from sqlalchemy import select

    from app.models import Article

    _with_catalog(db)
    article = db.scalar(select(Article).where(Article.name_norm == "zerkalo_100x200"))
    assert article is not None
    article.active = False
    db.commit()

    table = build_stock_table(db)
    assert table.row_count == 306
    row = next(row for row in table.rows if row.article_id == article.id)
    assert row.active is False
    assert row.stock_text(0) == "0"

    hidden = build_stock_table(db, hide_inactive=True)
    assert hidden.row_count == 305


def test_stock_counts_per_site(db: Session) -> None:
    """Остаток считается по площадкам."""
    _with_catalog(db)
    _make_stock(db, "zerkalo_100x200", 3, "Коминтерна")
    _make_stock(db, "zerkalo_100x200", 2, "Учительская")

    table = build_stock_table(db)
    index_by_site = {site.name: index for index, site in enumerate(table.sites)}
    row = next(row for row in table.rows if row.name == "Zerkalo_100х200")
    assert row.stock[index_by_site["Коминтерна"]] == 3
    assert row.stock[index_by_site["Учительская"]] == 2


def test_tasks_column_sums_plan(db: Session) -> None:
    """«Задачи на …» — сумма плана по площадке и дате; пусто, если задач нет (SPEC §12 п. 20)."""
    from sqlalchemy import select

    from app.models import Article

    _with_catalog(db)
    article = db.scalar(select(Article).where(Article.name_norm == "zerkalo_100x200"))
    assert article is not None
    lead = make_user(db, name="Начальник остатков", role=ROLE_SHIFT_LEAD)
    first = make_user(db, name="Сборщик остатков 1", role=ROLE_ASSEMBLER)
    second = make_user(db, name="Сборщик остатков 2", role=ROLE_ASSEMBLER)
    site = site_by_name(db)
    today = today_msk()

    for worker, qty in ((first, 5), (second, 7)):
        create_task(
            db,
            actor=lead,
            day=today,
            site_id=site.id,
            user_id=worker.id,
            kind="article",
            article_id=article.id,
            qty_plan=qty,
        )
    db.commit()

    table = build_stock_table(db, day=today)
    index = {s.name: i for i, s in enumerate(table.sites)}
    row = next(row for row in table.rows if row.name == "Zerkalo_100х200")
    assert row.tasks_text(index["Коминтерна"]) == "12"
    assert row.tasks_text(index["Учительская"]) == ""

    tomorrow = build_stock_table(db, day=today + timedelta(days=1))
    row_tomorrow = next(row for row in tomorrow.rows if row.name == "Zerkalo_100х200")
    assert row_tomorrow.tasks_text(index["Коминтерна"]) == ""


def test_tsv_has_one_value_per_row_including_headers(db: Session) -> None:
    """TSV содержит по строке на каждую строку таблицы, включая заголовки (SPEC §12 п. 19)."""
    _with_catalog(db)
    _make_stock(db, "zerkalo_100x200", 2)
    table = build_stock_table(db)

    tsv = table_to_tsv(table, [("stock", 0)])
    lines = tsv.split("\n")
    assert len(lines) == table.row_count == 306

    header_positions = [index for index, row in enumerate(table.rows) if row.is_header]
    for position in header_positions:
        assert lines[position] == ""

    first_article = next(index for index, row in enumerate(table.rows) if not row.is_header)
    assert lines[first_article] != ""


def test_tsv_two_columns_side_by_side(db: Session) -> None:
    """Кнопка «Копировать две площадки» даёт два столбца через табуляцию."""
    _with_catalog(db)
    _make_stock(db, "zerkalo_100x200", 1, "Учительская")
    table = build_stock_table(db)
    tsv = table_to_tsv(table, [("stock", 0), ("stock", 1)])
    lines = tsv.split("\n")
    assert len(lines) == 306
    assert all(line.count("\t") == 1 for line in lines)


def test_sheet_rows_match_screen(db: Session) -> None:
    """Лист «Остатки из системы» совпадает построчно с экраном (SPEC §12 п. 19)."""
    _with_catalog(db)
    _make_stock(db, "zerkalo_100x200", 4)
    table = build_stock_table(db)
    rows = sheet_rows(table)

    assert len(rows) == table.row_count
    assert [row[0] for row in rows] == [row.name for row in table.rows]
    for index, row in enumerate(table.rows):
        assert rows[index][1] == row.stock_text(0)
        assert rows[index][1 + len(table.sites)] == row.tasks_text(0)


def test_zero_cells_marked_on_page(client: TestClient, db: Session) -> None:
    """Нулевые остатки подсвечиваются, таблица открывается целиком."""
    _with_catalog(db)
    _make_stock(db, "zerkalo_100x200", 2)
    lead = make_user(db, name="Начальник экрана", role=ROLE_SHIFT_LEAD, pin="1234")
    login(client, lead)

    response = client.get("/stock")
    assert response.status_code == 200
    assert "ПОЛОТНА" in response.text
    assert 'class="num zero"' in response.text
    assert "Zerkalo_100х200" in response.text


def test_stock_exports(client: TestClient, db: Session) -> None:
    """CSV и XLSX отдаются с теми же строками."""
    _with_catalog(db)
    _make_stock(db, "zerkalo_100x200", 1)
    admin = make_user(db, name="Админ выгрузки", role=ROLE_ADMIN, pin="1234")
    login(client, admin)

    csv_response = client.get("/stock.csv")
    assert csv_response.status_code == 200
    text = csv_response.content.decode("utf-8")
    assert text.startswith("﻿")
    assert len(text.strip().split("\n")) == 307  # заголовок + 306 строк

    xlsx_response = client.get("/stock.xlsx")
    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )

    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(xlsx_response.content))
    sheet = workbook.active
    assert sheet.max_row == 307
    assert sheet.cell(2, 1).value == "ПОЛОТНА"


def test_loader_cannot_open_stock(client: TestClient, db: Session) -> None:
    """Грузчику экран остатков недоступен (SPEC §4)."""
    from app.statuses import ROLE_LOADER

    loader = make_user(db, name="Грузчик остатков", role=ROLE_LOADER, pin="1234")
    login(client, loader)
    assert client.get("/stock").status_code == 403
