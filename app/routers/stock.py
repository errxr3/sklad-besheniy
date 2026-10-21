"""Экран «Остатки» и выгрузки (SPEC §6.9, §10)."""

import csv
import io
import json
import logging
from datetime import date
from urllib.parse import quote, urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import Site
from app.services.sheets import is_configured, last_sync, minutes_since_sync, sync_stock_to_sheets
from app.services.shipment_report import CHANNEL_COLUMNS, build_shipment_table, shipment_xlsx
from app.services.stock import build_stock_table, site_locative
from app.timeutil import now_utc, today_msk
from app.web import GuardedRoute, page

logger = logging.getLogger(__name__)
router = APIRouter(route_class=GuardedRoute)

# Через сколько минут без выгрузки показывать плашку (SPEC §10)
SHEETS_STALE_MINUTES = 25


def _attachment(name: str, fallback: str) -> str:
    """Заголовок скачивания с русским именем файла (RFC 5987).

    Простое имя оставлено для старых браузеров, настоящее — в filename*.
    """
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quote(name)}"


@router.get("/shipments", response_class=HTMLResponse)
def shipments_page(request: Request, db: Session = Depends(get_db), day: str = "") -> Response:
    """Отгрузка за день — как гугл-таблица: остатки, задачи, каналы, «после отгрузок»."""
    try:
        target_day = date.fromisoformat(day) if day else today_msk()
    except ValueError:
        target_day = today_msk()
    table = build_shipment_table(db, day=target_day)
    return page(
        request,
        "shipments.html",
        {
            "table": table,
            "day": target_day,
            "channels": [title for title, _, _ in CHANNEL_COLUMNS],
        },
    )


@router.get("/shipments.xlsx")
def shipments_xlsx(request: Request, db: Session = Depends(get_db), day: str = "") -> Response:
    """Полный файл отгрузки за день в столбцах гугл-таблицы."""
    try:
        target_day = date.fromisoformat(day) if day else today_msk()
    except ValueError:
        target_day = today_msk()
    content = shipment_xlsx(build_shipment_table(db, day=target_day))
    name = f"Отгрузка_{target_day:%d.%m.%Y}_{now_utc().strftime('%H-%M')}.xlsx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _attachment(name, "shipments.xlsx")},
    )


@router.get("/stock", response_class=HTMLResponse)
def stock_page(
    request: Request,
    db: Session = Depends(get_db),
    day: str = "",
    hide_inactive: int = 0,
    ok: str = "",
    err: str = "",
) -> Response:
    """Таблица остатков в том же виде, что левый блок гугл-таблицы."""
    target_day = date.fromisoformat(day) if day else today_msk()
    table = build_stock_table(db, day=target_day, hide_inactive=bool(hide_inactive))

    columns = {
        "stock": [table.column_values("stock", index) for index in range(len(table.sites))],
        "tasks": [table.column_values("tasks", index) for index in range(len(table.sites))],
    }

    unmarked = list(
        db.scalars(
            select(Site).where(Site.active.is_(True), Site.labeling_done.is_(False))
        ).all()
    )
    stale_minutes = minutes_since_sync(db)
    sheets_warning = ""
    if is_configured():
        if stale_minutes is None:
            sheets_warning = "Гугл-таблица ещё ни разу не обновлялась"
        elif stale_minutes >= SHEETS_STALE_MINUTES:
            sheets_warning = f"Гугл-таблица не обновлялась {stale_minutes} минут"

    return page(
        request,
        "stock.html",
        {
            "table": table,
            "columns_json": json.dumps(columns, ensure_ascii=False),
            "day": target_day,
            "hide_inactive": bool(hide_inactive),
            "unmarked": unmarked,
            "sheets_configured": is_configured(),
            "sheets_warning": sheets_warning,
            "sheets_status": last_sync(db)[1],
            "ok": ok,
            "err": err,
        },
    )


@router.get("/stock.csv")
def stock_csv(request: Request, db: Session = Depends(get_db), day: str = "") -> Response:
    """Остатки в CSV: те же строки и тот же порядок, что на экране."""
    target_day = date.fromisoformat(day) if day else today_msk()
    table = build_stock_table(db, day=target_day)

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    header = ["Артикул"]
    header.extend(site.name for site in table.sites)
    header.extend(f"Задачи на {site_locative(site.name)}" for site in table.sites)
    writer.writerow(header)
    for row in table.rows:
        line = [row.name]
        line.extend(row.stock_text(index) for index in range(len(table.sites)))
        line.extend(row.tasks_text(index) for index in range(len(table.sites)))
        writer.writerow(line)

    name = f"Остатки_{target_day:%d.%m.%Y}_{now_utc().strftime('%H-%M')}.csv"
    return Response(
        content=("﻿" + buffer.getvalue()).encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": _attachment(name, "stock.csv")},
    )


@router.get("/stock.xlsx")
def stock_xlsx(request: Request, db: Session = Depends(get_db), day: str = "") -> Response:
    """Остатки в XLSX: те же строки, цвета подгрупп сохраняются."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    target_day = date.fromisoformat(day) if day else today_msk()
    table = build_stock_table(db, day=target_day)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Остатки"

    header = ["Артикул"]
    header.extend(site.name for site in table.sites)
    header.extend(f"Задачи на {site_locative(site.name)} ({target_day:%d.%m})" for site in table.sites)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    zero_fill = PatternFill("solid", fgColor="F8C9CD")
    for row in table.rows:
        values: list[str | int] = [row.name]
        for index in range(len(table.sites)):
            text = row.stock_text(index)
            values.append(int(text) if text else "")
        for index in range(len(table.sites)):
            text = row.tasks_text(index)
            values.append(int(text) if text else "")
        sheet.append(values)

        current = sheet.max_row
        if row.is_header:
            for cell in sheet[current]:
                cell.font = Font(bold=True)
        if row.color_hex:
            fill = PatternFill("solid", fgColor=row.color_hex.lstrip("#"))
            sheet.cell(current, 1).fill = fill
        if not row.is_header:
            for index in range(len(table.sites)):
                cell = sheet.cell(current, 2 + index)
                if cell.value == 0:
                    cell.fill = zero_fill

    sheet.column_dimensions["A"].width = 34
    sheet.freeze_panes = "B2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    name = f"Остатки_{target_day:%d.%m.%Y}_{now_utc().strftime('%H-%M')}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _attachment(name, "stock.xlsx")},
    )


@router.post("/stock/sync")
def stock_sync(request: Request, db: Session = Depends(get_db)) -> Response:
    """Кнопка «Выгрузить сейчас» (SPEC §10)."""
    result = sync_stock_to_sheets(db)
    params = {"ok": result.message} if result.ok else {"err": result.message}
    logger.info("Ручная выгрузка в гугл-таблицу", extra={"status": result.message})
    return RedirectResponse(f"/stock?{urlencode(params)}", status_code=303)
