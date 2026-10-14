"""Импорт справочника артикулов из TOP.xlsx (SPEC §9)."""

import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.articles import (
    apply_tint,
    format_hex,
    is_subgroup_header,
    nearest_palette_color,
    normalize_article,
    prefix_for,
)
from app.models import Article

# Листы файла = группы (SPEC §9)
SHEET_GROUPS: tuple[str, ...] = ("Полотна", "Алюминий", "Безрамные")

# Артикулы начинаются с 3-й строки: строка 1 — дата, строка 2 — заголовки столбцов
FIRST_DATA_ROW = 3

_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
# Порядок цветов в теме xlsx и порядок индексов темы в ячейках — разные:
# в файле clrScheme идёт dk1, lt1, dk2, lt2, accent1..6, hlink, folHlink,
# а индекс темы в ячейке: 0=lt1, 1=dk1, 2=lt2, 3=dk2, 4..9=accent1..6, 10=hlink, 11=folHlink.
_THEME_ORDER = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5",
                "accent6", "hlink", "folHlink")


@dataclass
class ParsedRow:
    """Строка справочника из файла."""

    kind: str  # 'article' | 'header'
    name: str
    name_norm: str | None
    group: str
    subgroup: str | None
    prefix: str | None
    color_hex: str | None
    sort_order: int


@dataclass
class ImportReport:
    """Итог импорта для отчёта на странице и в CLI (SPEC §9)."""

    added: int = 0
    existed: int = 0
    duplicates: int = 0
    skipped: int = 0
    headers: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def total_rows(self) -> int:
        """Всего строк справочника в файле (артикулы + заголовки подгрупп)."""
        return self.added + self.existed + self.headers

    def as_text(self) -> str:
        """Короткий текст отчёта."""
        return (
            f"Добавлено: {self.added} · уже было: {self.existed} · "
            f"заголовков подгрупп: {self.headers} · дублей: {self.duplicates} · "
            f"пропущено пустых: {self.skipped}"
        )


def _theme_colors(path: Path) -> list[tuple[int, int, int]]:
    """Цвета темы книги в порядке индексов темы."""
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("xl/theme/theme1.xml")
    except (KeyError, zipfile.BadZipFile):
        return []
    root = ElementTree.fromstring(xml)
    scheme = root.find(".//a:clrScheme", _NS)
    if scheme is None:
        return []
    by_name: dict[str, tuple[int, int, int]] = {}
    for child in scheme:
        name = child.tag.split("}", 1)[-1]
        srgb = child.find("a:srgbClr", _NS)
        sys_clr = child.find("a:sysClr", _NS)
        value = None
        if srgb is not None:
            value = srgb.get("val")
        elif sys_clr is not None:
            value = sys_clr.get("lastClr")
        if value:
            by_name[name] = (int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    return [by_name.get(name, (255, 255, 255)) for name in _THEME_ORDER]


def _cell_color(cell: Cell, theme: list[tuple[int, int, int]]) -> str | None:
    """Цвет заливки ячейки в виде #RRGGBB.

    Явный RGB берётся как есть; цвет темы разворачивается через палитру темы и
    оттенок, а затем заменяется ближайшим цветом палитры по умолчанию (SPEC §6.9).
    """
    fill = cell.fill
    if fill is None or fill.patternType != "solid":
        return None
    color = fill.fgColor
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str):
        if color.rgb in ("00000000", "FFFFFFFF"):
            return None
        rgb = (int(color.rgb[-6:-4], 16), int(color.rgb[-4:-2], 16), int(color.rgb[-2:], 16))
        return format_hex(rgb)
    if color.type == "theme" and theme:
        index = int(color.theme)
        if 0 <= index < len(theme):
            tinted = apply_tint(theme[index], float(color.tint or 0.0))
            return nearest_palette_color(tinted)
    return None


def parse_top(path: str | Path) -> tuple[list[ParsedRow], ImportReport]:
    """Разобрать TOP.xlsx в список строк справочника."""
    path = Path(path)
    report = ImportReport()
    theme = _theme_colors(path)
    workbook = load_workbook(path, data_only=True)
    rows: list[ParsedRow] = []
    seen: dict[str, str] = {}
    order = 0

    for sheet_name in SHEET_GROUPS:
        if sheet_name not in workbook.sheetnames:
            report.warnings.append(f"В файле нет листа «{sheet_name}»")
            continue
        sheet = workbook[sheet_name]
        subgroup: str | None = None
        for row_index in range(FIRST_DATA_ROW, sheet.max_row + 1):
            cell = sheet.cell(row_index, 1)
            raw = cell.value
            label = sheet.cell(row_index, 2).value
            if raw is None or not str(raw).strip():
                if label and str(label).strip():
                    subgroup = str(label).strip()
                report.skipped += 1
                continue
            name = str(raw).strip()
            if label and str(label).strip():
                subgroup = str(label).strip()
            order += 1
            if is_subgroup_header(name):
                rows.append(
                    ParsedRow(
                        kind="header",
                        name=name,
                        name_norm=None,
                        group=sheet_name,
                        subgroup=None,
                        prefix=None,
                        color_hex=None,
                        sort_order=order,
                    )
                )
                subgroup = name
                report.headers += 1
                continue
            norm = normalize_article(name)
            if norm in seen:
                report.duplicates += 1
                report.warnings.append(f"Дубль артикула «{name}» (уже есть «{seen[norm]}»)")
                order -= 1
                continue
            seen[norm] = name
            rows.append(
                ParsedRow(
                    kind="article",
                    name=name,
                    name_norm=norm,
                    group=sheet_name,
                    subgroup=subgroup,
                    prefix=prefix_for(name),
                    color_hex=_cell_color(cell, theme),
                    sort_order=order,
                )
            )
    workbook.close()
    return rows, report


def import_top(session: Session, path: str | Path) -> ImportReport:
    """Загрузить справочник в базу.

    Повторный импорт не создаёт дублей: артикулы ищутся по нормализованному имени,
    у найденных обновляются порядок, группа, подгруппа и цвет. Артикулы, которых нет
    в файле, не удаляются (их можно деактивировать вручную — SPEC §9).
    Строки-заголовки подгрупп пересоздаются целиком: на них ничего не ссылается.
    """
    rows, report = parse_top(path)

    existing = {a.name_norm: a for a in session.scalars(select(Article)).all() if a.name_norm}
    for old_header in session.scalars(select(Article).where(Article.kind == "header")).all():
        session.delete(old_header)
    session.flush()

    for row in rows:
        if row.kind == "header":
            session.add(
                Article(
                    kind="header",
                    name=row.name,
                    name_norm=None,
                    group=row.group,
                    subgroup=None,
                    prefix=None,
                    color_hex=None,
                    active=True,
                    sort_order=row.sort_order,
                )
            )
            continue
        article = existing.get(row.name_norm or "")
        if article is None:
            session.add(
                Article(
                    kind="article",
                    name=row.name,
                    name_norm=row.name_norm,
                    group=row.group,
                    subgroup=row.subgroup,
                    prefix=row.prefix,
                    color_hex=row.color_hex,
                    active=True,
                    sort_order=row.sort_order,
                )
            )
            report.added += 1
        else:
            article.name = row.name
            article.group = row.group
            article.subgroup = row.subgroup
            article.prefix = row.prefix
            article.color_hex = row.color_hex
            article.sort_order = row.sort_order
            report.existed += 1
    session.flush()
    return report
