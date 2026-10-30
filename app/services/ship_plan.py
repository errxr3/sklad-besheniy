"""Экран «План отгрузки» — утренняя таблица владельца, только заполненная сама (2026-10-30).

Строки — артикулы в порядке TOP.xlsx с заголовками подгрупп. По каждому артикулу:
остаток на складе по площадкам, сколько ещё в работе по задачам, сколько нужно на каждый
канал по заказам дня из BLAGO OTGRUZKI и сколько уже погружено, решение склада
(едет / ждёт ответа начальника / не едет).

Заказы дня — то, что BLAGO OTGRUZKI прислал в этот день (утренний список и волны).
Ozon утром присылается без машины, поэтому в плане Ozon одним столбцом,
а погружено в него — сумма по всем машинам Озон 1…4.
"""

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Article, ShipDemand, ShipQuestion
from app.services.batches import ASK, GO, LATER, NO, in_work_by_article_site
from app.services.shipment_report import CHANNEL_COLUMNS, shipped_by_article_column
from app.services.stock import active_sites, stock_at, stock_by_article_site
from app.statuses import OZON
from app.timeutil import fmt_dt, msk_day_bounds, now_utc, today_msk

# Столбцы плана: (заголовок, канал заказа)
PLAN_CHANNELS: tuple[tuple[str, str], ...] = (
    ("Озон", OZON),
    ("КГТ", "КГТ"),
    ("СДЭК", "СДЭК"),
    ("WB", "WB"),
    ("СГТ", "СГТ"),
)
_PLAN_INDEX = {channel: index for index, (_, channel) in enumerate(PLAN_CHANNELS)}


@dataclass
class PlanRow:
    """Строка плана: артикул или заголовок подгруппы."""

    name: str
    is_header: bool
    color_hex: str | None
    stock: list[int] = field(default_factory=list)  # по площадкам, как в шапке
    in_work: int = 0
    need: list[int] = field(default_factory=list)  # нужно по каналам
    loaded: list[int] = field(default_factory=list)  # погружено по каналам
    go: int = 0
    later: int = 0  # доделают к вечеру, поедет в Озон 3
    ask: int = 0
    no: int = 0
    no_reason: str = ""

    @property
    def need_total(self) -> int:
        return sum(self.need)

    @property
    def loaded_total(self) -> int:
        return sum(self.loaded)

    @property
    def stock_total(self) -> int:
        return sum(self.stock)

    @property
    def state(self) -> str:
        """Цвет строки: no — не хватит, ask — ждёт ответа, work — доделают к вечеру, ok, done."""
        if self.no:
            return "no"
        if self.ask:
            return "ask"
        if self.need_total and self.loaded_total >= self.need_total:
            return "done"
        if self.later or self.need_total > self.stock_total + self.loaded_total:
            return "work"
        return "ok"

    @property
    def active(self) -> bool:
        return bool(self.need_total or self.loaded_total)


@dataclass
class ShipPlan:
    """План отгрузки на день."""

    day: date
    site_names: list[str]
    rows: list[PlanRow]
    built_at: str
    columns: list[str] = field(default_factory=lambda: [title for title, _ in PLAN_CHANNELS])
    need: list[int] = field(default_factory=list)
    loaded: list[int] = field(default_factory=list)
    go: int = 0
    later: int = 0
    ask: int = 0
    no: int = 0
    orders: int = 0
    unknown: list[str] = field(default_factory=list)  # артикулов нет в справочнике склада
    open_questions: int = 0
    is_today: bool = True


def build_ship_plan(session: Session, *, day: date) -> ShipPlan:
    """Собрать план отгрузки на день."""
    start, end = msk_day_bounds(day)
    sites = active_sites(session)
    stock = stock_by_article_site(session) if day >= today_msk() else stock_at(session, end)
    work: dict[int, int] = defaultdict(int)
    for (article_id, _), qty in in_work_by_article_site(session, day=day).items():
        work[article_id] += qty

    lines = session.scalars(
        select(ShipDemand).where(ShipDemand.updated_at >= start, ShipDemand.updated_at < end)
    ).all()
    need: dict[int, list[int]] = defaultdict(lambda: [0] * len(PLAN_CHANNELS))
    decided: dict[int, Counter[str]] = defaultdict(Counter)
    reasons: dict[int, Counter[str]] = defaultdict(Counter)
    unknown: set[str] = set()
    for line in lines:
        if line.article_id is None:
            unknown.add(line.article_name)
            continue
        index = _PLAN_INDEX.get(line.channel)
        if index is not None:
            need[line.article_id][index] += line.qty
        decided[line.article_id][line.decision or ASK] += line.qty
        if line.decision == NO and line.reason:
            reasons[line.article_id][line.reason] += line.qty

    # Погружено за день по каналам: закрытые каналы и ещё открытые
    loaded: dict[int, list[int]] = defaultdict(lambda: [0] * len(PLAN_CHANNELS))
    for site in sites:
        pending: dict[tuple[int, int], int] = defaultdict(int)
        closed = shipped_by_article_column(session, day=day, site_id=site.id, pending=pending)
        for source in (closed, pending):
            for (article_id, column), count in source.items():
                if not 0 <= column < len(CHANNEL_COLUMNS):
                    continue
                index = _PLAN_INDEX.get(CHANNEL_COLUMNS[column][1])
                if index is not None:
                    loaded[article_id][index] += count

    rows: list[PlanRow] = []
    for article in session.scalars(select(Article).order_by(Article.sort_order, Article.id)).all():
        if article.kind == "header":
            rows.append(PlanRow(name=article.name, is_header=True, color_hex=article.color_hex))
            continue
        votes = decided.get(article.id, Counter())
        rows.append(
            PlanRow(
                name=article.name,
                is_header=False,
                color_hex=article.color_hex,
                stock=[stock.get((article.id, site.id), 0) for site in sites],
                in_work=work.get(article.id, 0),
                need=list(need.get(article.id, [0] * len(PLAN_CHANNELS))),
                loaded=list(loaded.get(article.id, [0] * len(PLAN_CHANNELS))),
                go=votes[GO],
                later=votes[LATER],
                ask=votes[ASK],
                no=votes[NO],
                no_reason=reasons[article.id].most_common(1)[0][0] if reasons.get(article.id) else "",
            )
        )
    items = [row for row in rows if not row.is_header]
    questions = session.scalars(
        select(ShipQuestion).where(ShipQuestion.day == day, ShipQuestion.answer == "open")
    ).all()
    return ShipPlan(
        day=day,
        site_names=[site.name for site in sites],
        rows=rows,
        built_at=fmt_dt(now_utc()),
        need=[sum(row.need[i] for row in items) for i in range(len(PLAN_CHANNELS))],
        loaded=[sum(row.loaded[i] for row in items) for i in range(len(PLAN_CHANNELS))],
        go=sum(row.go for row in items),
        later=sum(row.later for row in items),
        ask=sum(row.ask for row in items),
        no=sum(row.no for row in items),
        orders=len({(line.market, line.posting) for line in lines}),
        unknown=sorted(unknown),
        open_questions=sum(1 for q in questions if q.short > q.promised),
        is_today=day >= today_msk(),
    )


def visible_rows(plan: ShipPlan, *, show_all: bool) -> list[PlanRow]:
    """Строки для экрана: все или только артикулы, по которым есть заказы или отгрузка.

    Заголовок подгруппы показывается, если под ним есть хоть одна строка.
    """
    if show_all:
        return plan.rows
    result: list[PlanRow] = []
    header: PlanRow | None = None
    for row in plan.rows:
        if row.is_header:
            header = row
            continue
        if not row.active:
            continue
        if header is not None:
            result.append(header)
            header = None
        result.append(row)
    return result


STATE_LABELS = {
    "no": "не хватит",
    "ask": "ждёт ответа",
    "work": "доделают к вечеру",
    "ok": "хватает",
    "done": "погружено",
}


def ship_plan_xlsx(plan: ShipPlan) -> bytes:
    """План отгрузки в Excel: те же столбцы, что на экране, цвета решений."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    fills = {
        "no": PatternFill("solid", fgColor="F8D7DA"),
        "ask": PatternFill("solid", fgColor="FFF3CD"),
        "work": PatternFill("solid", fgColor="FFF3CD"),
        "done": PatternFill("solid", fgColor="D1E7DD"),
    }
    book = Workbook()
    sheet = book.active
    sheet.title = f"План {plan.day:%d.%m}"
    head = (
        ["Артикул", *[f"Склад {name}" for name in plan.site_names], "В работе"]
        + [f"{title} нужно" for title in plan.columns]
        + [f"{title} погружено" for title in plan.columns]
        + ["Едет", "Вечером (Озон 3)", "Ждёт ответа", "Не едет", "Итог", "Почему не едет"]
    )
    sheet.append(head)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in visible_rows(plan, show_all=False):
        if row.is_header:
            sheet.append([row.name])
            sheet.cell(sheet.max_row, 1).font = Font(bold=True)
            continue
        sheet.append(
            [row.name, *row.stock, row.in_work, *row.need, *row.loaded, row.go, row.later, row.ask, row.no,
             STATE_LABELS[row.state], row.no_reason]
        )
        fill = fills.get(row.state)
        if fill is not None:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    sheet.append(
        ["Всего", *[""] * len(plan.site_names), "", *plan.need, *plan.loaded,
         plan.go, plan.later, plan.ask, plan.no]
    )
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 32
    sheet.freeze_panes = "B2"
    out = BytesIO()
    book.save(out)
    return out.getvalue()
